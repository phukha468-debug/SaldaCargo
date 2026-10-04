# -*- coding: utf-8 -*-
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

DOCS_DIR = os.path.abspath("docs/reconciliation")
PUBLIC_DIR = os.path.abspath("apps/web/public/documents")
os.makedirs(DOCS_DIR, exist_ok=True)
os.makedirs(PUBLIC_DIR, exist_ok=True)

XLSX_PATH = os.path.join(DOCS_DIR, "akt-sverki-uvs-2026.xlsx")
HTML_PATH = os.path.join(DOCS_DIR, "akt-sverki-uvs-2026.html")
PDF_PATH = os.path.join(DOCS_DIR, "akt-sverki-uvs-2026.pdf")
PNG_PATH = os.path.join(DOCS_DIR, "akt-sverki-uvs-2026.png")

# Чистые данные без примечаний и без расшифровок
operations = [
    # (Дата, Документ / Основание, Рейс, Водитель, Авто, Дебет, Кредит)
    ("18.05.2026", "Счёт №1 от 18.05.2026", "Рейс №88", "Никифоров Д.В.", "Газель Next", 7000, 0),
    ("18.05.2026", "Оплата по QR от 18.05.2026", "Рейс №88", "—", "—", 0, 7000),
    ("25.05.2026", "Счёт №2 от 25.05.2026", "Рейс №130", "Никифоров Д.В.", "Газель Next", 6000, 0),
    ("25.05.2026", "Оплата по QR от 25.05.2026", "Рейс №130", "—", "—", 0, 6000),
    ("18.06.2026", "Счёт №3 от 18.06.2026", "Рейс №249", "Никифоров Д.В.", "Газель Next", 5000, 0),
    ("19.06.2026", "Рейс №255", "Рейс №255", "Никифоров Д.В.", "Газель Next", 3000, 0),
    ("21.06.2026", "Рейс №259", "Рейс №259", "Никифоров Д.В.", "Газель Next", 6000, 0),
    ("22.06.2026", "Рейс №264", "Рейс №264", "Никифоров Д.В.", "Газель Next", 3000, 0),
    ("24.06.2026", "Рейс №275", "Рейс №275", "Никифоров Д.В.", "Газель Next", 6000, 0),
    ("25.06.2026", "Рейс №283", "Рейс №283", "Никифоров Д.В.", "Газель Next", 3000, 0),
    ("25.06.2026", "Счёт №15 от 25.06.2026", "Рейс №286", "Никифоров Д.В.", "Газель Next", 8000, 0),
    ("30.06.2026", "Рейс №293", "Рейс №293", "Никифоров Д.В.", "Газель Next", 4000, 0),
    ("01.07.2026", "Платёж б/н от 01.07.2026", "—", "—", "—", 0, 3000),
    ("01.07.2026", "Платёж б/н от 01.07.2026", "—", "—", "—", 0, 3000),
    ("07.07.2026", "Счёт №10 от 13.07.2026", "Рейс №319", "Никифоров Д.В.", "Газель Next", 12000, 0),
    ("08.07.2026", "Счёт №11 от 13.07.2026", "Рейс №322", "Никифоров Д.В.", "Газель Next", 4000, 0),
    ("09.07.2026", "Счёт №12 от 12.07.2026", "Рейс №329", "Никифоров Д.В.", "Газель Next", 6000, 0),
    ("10.07.2026", "Счёт №13 от 13.07.2026", "Рейс №334", "Никифоров Д.В.", "Газель Next", 3000, 0),
    ("14.07.2026", "Счёт №14 от 26.08.2026", "Рейс №350", "Никифоров Д.В.", "Газель Next", 6000, 0),
    ("25.08.2026", "Счёт №16 от 27.09.2026", "Рейс №566", "Никифоров Д.В.", "Газель Next", 6000, 0),
    ("27.08.2026", "Счёт №17 от 27.08.2026", "Рейс №578", "Полещук В.В.", "Газель", 5000, 0),
    ("04.09.2026", "Платёж б/н от 04.09.2026", "—", "—", "—", 0, 24000),
    ("08.09.2026", "Счёт №18 от 10.09.2026", "Рейс №639", "Фарух", "Газель Next", 8000, 0),
    ("29.09.2026", "Платёж б/н от 29.09.2026", "—", "—", "—", 0, 15000),
    ("02.10.2026", "Счёт №19 от 03.10.2026", "Рейс №756", "Иван Фирсов", "Газель Next", 5000, 0),
]

unpaid_orders = [
    # (№, Дата, Счёт / Акт, Рейс, Водитель, Автомобиль, Сумма)
    (1, "25.06.2026", "Счёт №15 от 25.06.2026", "Рейс №286", "Никифоров Д.В.", "Газель Next", 8000),
    (2, "08.07.2026", "Счёт №11 от 13.07.2026", "Рейс №322", "Никифоров Д.В.", "Газель Next", 4000),
    (3, "09.07.2026", "Счёт №12 от 12.07.2026", "Рейс №329", "Никифоров Д.В.", "Газель Next", 6000),
    (4, "14.07.2026", "Счёт №14 от 26.08.2026", "Рейс №350", "Никифоров Д.В.", "Газель Next", 6000),
    (5, "25.08.2026", "Счёт №16 от 27.09.2026", "Рейс №566", "Никифоров Д.В.", "Газель Next", 6000),
    (6, "27.08.2026", "Счёт №17 от 27.08.2026", "Рейс №578", "Полещук В.В.", "Газель", 5000),
    (7, "08.09.2026", "Счёт №18 от 10.09.2026", "Рейс №639", "Фарух", "Газель Next", 8000),
    (8, "02.10.2026", "Счёт №19 от 03.10.2026", "Рейс №756", "Иван Фирсов", "Газель Next", 5000),
]

def create_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Акт сверки УВС"

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    header_border = Border(
        left=Side(style='thin', color='0F172A'),
        right=Side(style='thin', color='0F172A'),
        top=Side(style='medium', color='0F172A'),
        bottom=Side(style='medium', color='0F172A')
    )

    # Title
    ws.merge_cells("A1:H1")
    ws["A1"] = "АКТ СВЕРКИ ВЗАИМНЫХ РАСЧЕТОВ"
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color="0F172A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:H2")
    ws["A2"] = "за период: 18 мая 2026 г. — 04 октября 2026 г."
    ws["A2"].font = Font(name="Calibri", size=11, italic=True, color="475569")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Parties
    ws.merge_cells("A4:D4")
    ws["A4"] = "Исполнитель: ИП Нигамедьянов А.С. (ТК501 / SaldaCargo)"
    ws["A4"].font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")

    ws.merge_cells("E4:H4")
    ws["E4"] = "Заказчик: ООО «Уралводстрой» (УВС)"
    ws["E4"].font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")

    # Table Header
    headers = [
        "№", "Дата", "Документ / Основание", "Рейс", "Водитель", "Автомобиль", 
        "Дебет (Оказано), ₽", "Кредит (Оплачено), ₽"
    ]
    header_row = 6
    ws.row_dimensions[header_row].height = 22
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    # Data rows
    cur_row = 7
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, op in enumerate(operations, 1):
        date_str, doc_str, trip_str, drv_str, car_str, debit, credit = op
        row_fill = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)

        cells = [
            (idx, "center"),
            (date_str, "center"),
            (doc_str, "left"),
            (trip_str, "center"),
            (drv_str, "left"),
            (car_str, "left"),
            (debit, "right"),
            (credit, "right")
        ]

        for col_idx, (val, align) in enumerate(cells, 1):
            c = ws.cell(row=cur_row, column=col_idx, value=val)
            c.border = thin_border
            if row_fill.fill_type:
                c.fill = row_fill
            c.alignment = Alignment(horizontal=align, vertical="center")
            if col_idx in [7, 8]:
                c.number_format = '#,##0.00 "₽"'
                c.font = Font(name="Calibri", size=10, bold=(val > 0))
            else:
                c.font = Font(name="Calibri", size=10)

        cur_row += 1

    # Total Row
    total_row = cur_row
    ws.merge_cells(f"A{total_row}:F{total_row}")
    ws[f"A{total_row}"] = "ИТОГО ОБОРОТЫ ЗА ПЕРИОД:"
    ws[f"A{total_row}"].font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    ws[f"A{total_row}"].alignment = Alignment(horizontal="right", vertical="center")

    c_deb = ws.cell(row=total_row, column=7, value=f"=SUM(G7:G{total_row-1})")
    c_deb.number_format = '#,##0.00 "₽"'
    c_deb.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    c_deb.alignment = Alignment(horizontal="right", vertical="center")

    c_cred = ws.cell(row=total_row, column=8, value=f"=SUM(H7:H{total_row-1})")
    c_cred.number_format = '#,##0.00 "₽"'
    c_cred.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    c_cred.alignment = Alignment(horizontal="right", vertical="center")

    tot_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    for col in range(1, 9):
        c = ws.cell(row=total_row, column=col)
        c.fill = tot_fill
        c.border = header_border

    # Balance Row
    saldo_row = total_row + 1
    ws.merge_cells(f"A{saldo_row}:F{saldo_row}")
    ws[f"A{saldo_row}"] = "КОНЕЧНОЕ САЛЬДО (ЗАДОЛЖЕННОСТЬ УВС):"
    ws[f"A{saldo_row}"].font = Font(name="Calibri", size=11, bold=True, color="B91C1C")
    ws[f"A{saldo_row}"].alignment = Alignment(horizontal="right", vertical="center")

    c_saldo = ws.cell(row=saldo_row, column=7, value=f"=G{total_row}-H{total_row}")
    c_saldo.number_format = '#,##0.00 "₽"'
    c_saldo.font = Font(name="Calibri", size=11, bold=True, color="B91C1C")
    c_saldo.alignment = Alignment(horizontal="right", vertical="center")

    c_saldo_lbl = ws.cell(row=saldo_row, column=8, value="В пользу ИП Нигамедьянов")
    c_saldo_lbl.font = Font(name="Calibri", size=9, bold=True, color="B91C1C")
    c_saldo_lbl.alignment = Alignment(horizontal="center", vertical="center")

    saldo_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    for col in range(1, 9):
        c = ws.cell(row=saldo_row, column=col)
        c.fill = saldo_fill
        c.border = header_border

    # Section for UNPAID ORDERS
    cur_row = saldo_row + 3
    ws.merge_cells(f"A{cur_row}:G{cur_row}")
    ws[f"A{cur_row}"] = "РЕЕСТР НЕОПЛАЧЕННЫХ РЕЙСОВ"
    ws[f"A{cur_row}"].font = Font(name="Calibri", size=11, bold=True, color="991B1B")
    ws[f"A{cur_row}"].alignment = Alignment(horizontal="left", vertical="center")

    cur_row += 1
    unpaid_headers = ["№", "Дата рейса", "Счёт / Акт", "Рейс", "Водитель", "Автомобиль", "Сумма к оплате, ₽"]
    for col_idx, h in enumerate(unpaid_headers, 1):
        c = ws.cell(row=cur_row, column=col_idx, value=h)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="7F1D1D", end_color="7F1D1D", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = header_border

    unpaid_start = cur_row + 1
    for row_data in unpaid_orders:
        cur_row += 1
        num, d_str, inv_str, r_str, drv_str, car_str, amt = row_data
        row_vals = [num, d_str, inv_str, r_str, drv_str, car_str, amt]
        for col_idx, v in enumerate(row_vals, 1):
            c = ws.cell(row=cur_row, column=col_idx, value=v)
            c.border = thin_border
            c.fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
            if col_idx in [1, 2, 4]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 7:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0.00 "₽"'
                c.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
            c.font = Font(name="Calibri", size=9)

    # Unpaid total
    cur_row += 1
    ws.merge_cells(f"A{cur_row}:F{cur_row}")
    ws[f"A{cur_row}"] = "ИТОГО К ОПЛАТЕ:"
    ws[f"A{cur_row}"].font = Font(name="Calibri", size=10, bold=True, color="991B1B")
    ws[f"A{cur_row}"].alignment = Alignment(horizontal="right", vertical="center")

    c_unpaid_tot = ws.cell(row=cur_row, column=7, value=f"=SUM(G{unpaid_start}:G{cur_row-1})")
    c_unpaid_tot.number_format = '#,##0.00 "₽"'
    c_unpaid_tot.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
    c_unpaid_tot.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, 8):
        c = ws.cell(row=cur_row, column=col)
        c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        c.border = header_border

    # Signatures
    cur_row += 3
    ws.merge_cells(f"A{cur_row}:C{cur_row}")
    ws[f"A{cur_row}"] = "От Исполнителя:"
    ws[f"A{cur_row}"].font = Font(name="Calibri", size=10, bold=True)

    ws.merge_cells(f"E{cur_row}:G{cur_row}")
    ws[f"E{cur_row}"] = "От Заказчика:"
    ws[f"E{cur_row}"].font = Font(name="Calibri", size=10, bold=True)

    cur_row += 1
    ws.merge_cells(f"A{cur_row}:C{cur_row}")
    ws[f"A{cur_row}"] = "ИП Нигамедьянов А.С."

    ws.merge_cells(f"E{cur_row}:G{cur_row}")
    ws[f"E{cur_row}"] = "ООО «Уралводстрой» (УВС)"

    cur_row += 2
    ws.merge_cells(f"A{cur_row}:C{cur_row}")
    ws[f"A{cur_row}"] = "Подпись: _________________ / Нигамедьянов А.С. /"

    ws.merge_cells(f"E{cur_row}:G{cur_row}")
    ws[f"E{cur_row}"] = "Подпись: _________________ / _________________ /"

    cur_row += 1
    ws.merge_cells(f"A{cur_row}:C{cur_row}")
    ws[f"A{cur_row}"] = "М.П."

    ws.merge_cells(f"E{cur_row}:G{cur_row}")
    ws[f"E{cur_row}"] = "М.П."

    # Column widths
    col_widths = [5, 12, 28, 14, 20, 18, 18, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(XLSX_PATH)
    print("Excel file created at:", XLSX_PATH)

def generate_html_and_pdf():
    rows_html = ""
    for idx, op in enumerate(operations, 1):
        date_str, doc_str, trip_str, drv_str, car_str, debit, credit = op
        bg = "#f8fafc" if idx % 2 == 0 else "#ffffff"

        deb_str = f"{debit:,.2f} ₽".replace(",", " ") if debit > 0 else "—"
        cred_str = f"{credit:,.2f} ₽".replace(",", " ") if credit > 0 else "—"

        rows_html += f"""
        <tr style="background: {bg};">
          <td style="text-align: center; color: #64748b;">{idx}</td>
          <td style="text-align: center; white-space: nowrap;">{date_str}</td>
          <td><b>{doc_str}</b></td>
          <td style="text-align: center; white-space: nowrap;">{trip_str}</td>
          <td>{drv_str}</td>
          <td>{car_str}</td>
          <td style="text-align: right; font-weight: bold; color: #0f172a;">{deb_str}</td>
          <td style="text-align: right; font-weight: bold; color: #15803d;">{cred_str}</td>
        </tr>
        """

    unpaid_rows_html = ""
    for row in unpaid_orders:
        num, d_str, inv_str, r_str, drv_str, car_str, amt = row
        amt_str = f"{amt:,.2f} ₽".replace(",", " ")
        unpaid_rows_html += f"""
        <tr style="background: #fff1f2;">
          <td style="text-align: center; font-weight: bold; color: #991b1b;">{num}</td>
          <td style="text-align: center; white-space: nowrap;">{d_str}</td>
          <td><b>{inv_str}</b></td>
          <td style="text-align: center; white-space: nowrap;">{r_str}</td>
          <td>{drv_str}</td>
          <td>{car_str}</td>
          <td style="text-align: right; font-weight: 900; color: #b91c1c; font-size: 13px;">{amt_str}</td>
        </tr>
        """

    html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<title>Акт сверки взаимных расчетов — УВС (2026)</title>
<style>
  @page {{
    size: A4 portrait;
    margin: 10mm 10mm 10mm 10mm;
  }}
  * {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    color: #0f172a;
    background: #ffffff;
    font-size: 11px;
    line-height: 1.35;
  }}
  .header {{
    border-bottom: 2px solid #0f172a;
    padding-bottom: 6px;
    margin-bottom: 10px;
  }}
  .title {{
    font-size: 17px;
    font-weight: 900;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    color: #0f172a;
    text-align: center;
  }}
  .subtitle {{
    font-size: 11px;
    font-weight: 600;
    color: #475569;
    text-align: center;
    margin-top: 2px;
  }}
  .parties-grid {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 16px;
    margin-bottom: 10px;
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 6px;
    padding: 8px 12px;
  }}
  .party-box h3 {{
    font-size: 11px;
    font-weight: 800;
    color: #1e3a8a;
    margin-bottom: 2px;
  }}
  .party-box p {{
    font-size: 10px;
    color: #334155;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin-bottom: 10px;
    font-size: 9.5px;
  }}
  th {{
    background: #1e293b;
    color: #ffffff;
    padding: 5px 6px;
    text-align: left;
    font-weight: 800;
    border: 1px solid #334155;
    font-size: 9px;
  }}
  td {{
    padding: 4px 6px;
    border: 1px solid #cbd5e1;
    line-height: 1.25;
  }}
  .summary-box {{
    background: #f1f5f9;
    border: 1.5px solid #cbd5e1;
    border-radius: 6px;
    padding: 8px 12px;
    margin-bottom: 10px;
  }}
  .summary-grid {{
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 8px;
    text-align: center;
  }}
  .summary-item .label {{
    font-size: 9.5px;
    font-weight: 700;
    color: #64748b;
    text-transform: uppercase;
  }}
  .summary-item .val {{
    font-size: 14px;
    font-weight: 900;
    margin-top: 1px;
  }}
  .val-deb {{ color: #0f172a; }}
  .val-cred {{ color: #15803d; }}
  .val-saldo {{ color: #b91c1c; }}
  .conclusion {{
    margin-top: 6px;
    padding-top: 5px;
    border-top: 1px dashed #94a3b8;
    font-size: 10.5px;
    font-weight: 700;
    color: #0f172a;
    text-align: center;
  }}
  .conclusion span {{
    color: #b91c1c;
    font-weight: 900;
  }}
  .unpaid-section {{
    margin-top: 10px;
    margin-bottom: 12px;
    page-break-inside: avoid;
  }}
  .unpaid-title {{
    font-size: 11px;
    font-weight: 900;
    color: #991b1b;
    margin-bottom: 5px;
    text-transform: uppercase;
  }}
  .signatures-grid {{
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 30px;
    margin-top: 15px;
    page-break-inside: avoid;
  }}
  .sign-col h4 {{
    font-size: 10.5px;
    font-weight: 800;
    color: #0f172a;
    border-bottom: 1px solid #cbd5e1;
    padding-bottom: 3px;
    margin-bottom: 10px;
  }}
  .sign-line {{
    margin-top: 25px;
    border-bottom: 1px solid #475569;
    width: 85%;
    display: flex;
    justify-content: space-between;
    font-size: 8.5px;
    color: #64748b;
    padding-bottom: 2px;
  }}
  .stamp-place {{
    margin-top: 10px;
    font-size: 9.5px;
    font-weight: bold;
    color: #64748b;
  }}
</style>
</head>
<body>

  <div class="header">
    <div class="title">Акт сверки взаимных расчетов</div>
    <div class="subtitle">за период: с 18 мая 2026 г. по 04 октября 2026 г.</div>
  </div>

  <div class="parties-grid">
    <div class="party-box">
      <h3>Исполнитель:</h3>
      <p><b>ИП Нигамедьянов А.С. (ТК501 / SaldaCargo)</b></p>
    </div>
    <div class="party-box">
      <h3>Заказчик:</h3>
      <p><b>ООО «Уралводстрой» (УВС)</b></p>
    </div>
  </div>

  <!-- Таблица операций -->
  <table>
    <thead>
      <tr>
        <th style="width: 20px; text-align: center;">№</th>
        <th style="width: 65px; text-align: center;">Дата</th>
        <th>Документ / Основание</th>
        <th style="width: 70px; text-align: center;">Рейс</th>
        <th style="width: 120px;">Водитель</th>
        <th style="width: 100px;">Автомобиль</th>
        <th style="width: 85px; text-align: right;">Оказано, ₽</th>
        <th style="width: 85px; text-align: right;">Оплачено, ₽</th>
      </tr>
    </thead>
    <tbody>
      {rows_html}
    </tbody>
  </table>

  <!-- Сводный блок расчетов -->
  <div class="summary-box">
    <div class="summary-grid">
      <div class="summary-item">
        <div class="label">Всего оказано услуг (Дебет)</div>
        <div class="val val-deb">106 000,00 ₽</div>
      </div>
      <div class="summary-item">
        <div class="label">Всего оплачено (Кредит)</div>
        <div class="val val-cred">58 000,00 ₽</div>
      </div>
      <div class="summary-item">
        <div class="label">Задолженность УВС (Сальдо)</div>
        <div class="val val-saldo">48 000,00 ₽</div>
      </div>
    </div>
    <div class="conclusion">
      Задолженность ООО «Уралводстрой» (УВС) в пользу ИП Нигамедьянов А.С. составляет 
      <span>48 000,00 (Сорок восемь тысяч) рублей 00 копеек</span>.
    </div>
  </div>

  <!-- Реестр неоплаченных рейсов -->
  <div class="unpaid-section">
    <div class="unpaid-title">Реестр неоплаченных рейсов:</div>
    <table>
      <thead>
        <tr style="background: #991b1b;">
          <th style="width: 20px; text-align: center;">№</th>
          <th style="width: 70px; text-align: center;">Дата</th>
          <th>Счёт / Акт</th>
          <th style="width: 70px; text-align: center;">Рейс</th>
          <th>Водитель</th>
          <th>Автомобиль</th>
          <th style="width: 95px; text-align: right;">Сумма, ₽</th>
        </tr>
      </thead>
      <tbody>
        {unpaid_rows_html}
        <tr style="background: #fee2e2; font-weight: 900;">
          <td colspan="6" style="text-align: right; font-size: 10px; color: #991b1b;">ИТОГО К ОПЛАТЕ:</td>
          <td style="text-align: right; color: #991b1b; font-size: 12px;">48 000,00 ₽</td>
        </tr>
      </tbody>
    </table>
  </div>

  <!-- Подписи сторон -->
  <div class="signatures-grid">
    <div class="sign-col">
      <h4>От Исполнителя: ИП Нигамедьянов А.С.</h4>
      <div class="sign-line">
        <span>(подпись)</span>
        <span>/ Нигамедьянов А.С. /</span>
      </div>
      <div class="stamp-place">М. П.</div>
    </div>
    <div class="sign-col">
      <h4>От Заказчика: ООО «Уралводстрой» (УВС)</h4>
      <div class="sign-line">
        <span>(подпись)</span>
        <span>/ ___________________________ /</span>
      </div>
      <div class="stamp-place">М. П.</div>
    </div>
  </div>

</body>
</html>"""

    with open(HTML_PATH, "w", encoding="utf-8") as f:
        f.write(html)
    print("HTML created at:", HTML_PATH)

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.set_content(html, wait_until="networkidle")
        page.pdf(
            path=PDF_PATH,
            format="A4",
            print_background=True,
            margin={"top": "8mm", "bottom": "8mm", "left": "8mm", "right": "8mm"}
        )
        page.screenshot(path=PNG_PATH, full_page=True)
        browser.close()
    print("PDF created at:", PDF_PATH)
    print("PNG created at:", PNG_PATH)

if __name__ == "__main__":
    create_excel()
    generate_html_and_pdf()
